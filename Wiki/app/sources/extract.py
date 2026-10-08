"""Turn uploaded files / URLs into plain text + sections for note extraction."""
import io
import re
from dataclasses import dataclass

import httpx
import pymupdf
import trafilatura

from . import pdf_ifrs
from .text import clean, split_long

TEXT_EXT = (".txt", ".md", ".markdown")
SUPPORTED_EXT = (".pdf", ".docx") + TEXT_EXT


@dataclass
class Extracted:
    title: str
    kind: str
    text: str
    sections: list[str]
    meta: dict


def _pack(blocks: list[str], max_words: int) -> list[str]:
    """Pack text blocks into sections of roughly max_words, never splitting a block unless it alone is too long."""
    sections, buf, n = [], [], 0
    for b in blocks:
        for piece in split_long(b, max_words):
            w = len(piece.split())
            if buf and n + w > max_words:
                sections.append("\n\n".join(buf))
                buf, n = [], 0
            buf.append(piece)
            n += w
    if buf:
        sections.append("\n\n".join(buf))
    return sections


def _blocks(text: str) -> list[str]:
    return [b.strip() for b in re.split(r"\n\s*\n", text) if b.strip()]


def _title_from(name: str) -> str:
    return re.sub(r"\.[a-z0-9]+$", "", name, flags=re.I).replace("_", " ").strip()


def from_pdf(name: str, data: bytes, max_words: int) -> Extracted:
    doc = pymupdf.open(stream=data, filetype="pdf")
    meta = {"pages": doc.page_count, **{k: v for k, v in (doc.metadata or {}).items() if v}}
    if pdf_ifrs.looks_like_ifrs(doc):
        # Binding text only, labelled by paragraph number so notes can cite it
        chunks = [c for c in pdf_ifrs.parse(doc) if c.source_type in ("standard", "appendix_a", "appendix_b")]
        blocks = [f"[{c.para_ref or 'p' + str(c.page)}] {c.content}" for c in chunks]
        meta["parser"] = "ifrs"
    else:
        blocks = [f"[page {i}]\n{p.get_text()}" for i, p in enumerate(doc, start=1) if p.get_text().strip()]
        meta["parser"] = "pdf"
    text = "\n\n".join(blocks)
    return Extracted(_title_from(name), "pdf", text, _pack(blocks, max_words), meta)


def from_docx(name: str, data: bytes, max_words: int) -> Extracted:
    import docx  # python-docx

    d = docx.Document(io.BytesIO(data))
    blocks = []
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        blocks.append(f"## {t}" if p.style.name.lower().startswith("heading") else t)
    for table in d.tables:
        rows = [" | ".join(clean(c.text) for c in r.cells) for r in table.rows]
        blocks.append("\n".join(rows))
    text = "\n\n".join(blocks)
    title = d.core_properties.title or _title_from(name)
    return Extracted(title, "docx", text, _pack(blocks, max_words), {})


def from_text(name: str, data: bytes, max_words: int, title: str | None = None) -> Extracted:
    text = data.decode("utf-8", errors="replace").replace("\r\n", "\n")
    kind = "markdown" if name.lower().endswith((".md", ".markdown")) else "text"
    if not title and kind == "markdown":
        m = re.search(r"^#\s+(.+)$", text, re.M)
        title = m.group(1).strip() if m else None
    return Extracted(title or _title_from(name), kind, text, _pack(_blocks(text), max_words), {})


async def from_url(url: str, max_words: int) -> Extracted:
    async with httpx.AsyncClient(follow_redirects=True, timeout=30, headers={"User-Agent": "Mozilla/5.0 (compatible; llm-wiki/0.1; +https://github.com/)", "Accept": "text/html,application/pdf;q=0.9,*/*;q=0.8"}) as client:
        resp = await client.get(url)
        resp.raise_for_status()
    ctype = resp.headers.get("content-type", "")
    if "pdf" in ctype or url.lower().endswith(".pdf"):
        ex = from_pdf(url.rsplit("/", 1)[-1] or "document.pdf", resp.content, max_words)
        ex.kind = "url"
        return ex
    html = resp.text
    text = trafilatura.extract(html, output_format="markdown", include_tables=True, include_links=False) or ""
    if not text.strip():
        raise ValueError("Could not extract readable text from URL")
    meta = trafilatura.extract_metadata(html)
    title = (meta.title if meta and meta.title else None) or url
    return Extracted(title, "url", text, _pack(_blocks(text), max_words), {"url": url, "content_type": ctype})


def from_upload(name: str, data: bytes, max_words: int) -> Extracted:
    low = name.lower()
    if low.endswith(".pdf"):
        return from_pdf(name, data, max_words)
    if low.endswith(".docx"):
        return from_docx(name, data, max_words)
    if low.endswith(TEXT_EXT):
        return from_text(name, data, max_words)
    raise ValueError(f"Unsupported file type. Supported: {', '.join(SUPPORTED_EXT)}")
