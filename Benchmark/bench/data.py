"""Load benchmark topics (query set) and human reference answers (assessment docx). Run: python -m bench.data"""
import difflib
import re
import sys
from dataclasses import dataclass, field

import docx

from .config import get_settings

TOPIC_RE = re.compile(r"^TOPIC (\d+) of (\d+):\s*(.+)$", re.M)


@dataclass
class Topic:
    number: int
    code: str  # e.g. "2.1.2"
    title: str  # e.g. "Expected Lifetime"
    questions: list[str] = field(default_factory=list)
    reference_label: str = ""
    reference: str = ""
    docx_title: str = ""
    section: str = ""

    def question_text(self) -> str:
        return f"{self.title}\n" + "\n".join(self.questions)


def load_topics(path) -> list[Topic]:
    text = path.read_text(encoding="utf-8").replace("\r\n", "\n")
    matches = list(TOPIC_RE.finditer(text))
    topics = []
    for i, m in enumerate(matches):
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        block = text[m.start() : end]
        code, _, title = m.group(3).strip().partition(" ")
        topics.append(Topic(int(m.group(1)), code, title.strip(), re.findall(r"^Q\d+\..*$", block, re.M)))
    return topics


def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def load_references(path) -> list[dict]:
    """Heading-3 topic titles in document order, paired with tables 1..N (table 0 is the status summary)."""
    d = docx.Document(str(path))
    titles, sections, section = [], [], ""
    for p in d.paragraphs:
        if p.style.name == "Heading 2":
            section = p.text.strip()
        elif p.style.name == "Heading 3" and p.text.strip():
            titles.append(p.text.strip())
            sections.append(section)
    tables = d.tables[1:]
    if len(tables) != len(titles):
        raise ValueError(f"{len(titles)} topic headings but {len(tables)} answer tables")
    refs = []
    for title, sec, t in zip(titles, sections, tables):
        first = t.rows[0].cells
        refs.append({"title": title, "section": sec, "label": first[0].text.strip(), "text": first[1].text.strip()})
    return refs


def load() -> list[Topic]:
    s = get_settings()
    topics = load_topics(s.query_set)
    refs = load_references(s.reference_docx)
    if len(topics) != len(refs):
        raise ValueError(f"{len(topics)} topics vs {len(refs)} references")
    for t, r in zip(topics, refs):
        ratio = difflib.SequenceMatcher(None, _norm(t.title), _norm(r["title"])).ratio()
        if ratio < 0.6 and _norm(t.title) not in _norm(r["title"]):
            raise ValueError(f"Topic {t.number} '{t.title}' does not match docx heading '{r['title']}' ({ratio:.2f})")
        if not r["text"]:
            raise ValueError(f"Topic {t.number}: empty reference answer")
        t.reference_label, t.reference, t.docx_title, t.section = r["label"], r["text"], r["title"], r["section"]
    return topics


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    for t in load():
        print(f"{t.number:>2} {t.code:<6} {t.title[:55]:<55} | {t.docx_title[:45]:<45} | {t.reference_label[:22]:<22} | "
              f"{len(t.reference.split())}w, {len(t.questions)}Q")
